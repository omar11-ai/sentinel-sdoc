"""Document parsing lanes: txt / pdf / xlsx / docx -> uniform lines + health status."""
from __future__ import annotations

import io

from pathlib import Path

# status: ok | empty | corrupt | image_only
def to_lines(path: Path) -> tuple[list[str], str]:
    try:
        data = path.read_bytes()
    except Exception:
        return [], "corrupt"
    if not data.strip():
        return [], "empty"

    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _pdf_lines(data)
    if suffix in (".xlsx", ".xls"):
        return _xlsx_lines(data)
    if suffix in (".docx", ".doc"):
        return _docx_lines(data)
    text = data.decode("utf-8", errors="replace")
    if data:
        ctrl = sum(1 for b in data if b < 9 or (13 < b < 32))
        if ctrl / max(len(data), 1) > 0.05:
            return [], "corrupt"
    return text.splitlines(), "ok"


def _pdf_lines(data: bytes) -> tuple[list[str], str]:
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        chunks: list[str] = []
        for page in reader.pages:
            try:
                chunks.append(page.extract_text() or "")
            except Exception:
                chunks.append("")
        text = "\n".join(chunks).strip()
        if not text:
            return [], "image_only"
        return text.splitlines(), "ok"
    except Exception:
        return [], "corrupt"


def _xlsx_lines(data: bytes) -> tuple[list[str], str]:
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        lines: list[str] = []
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                raw = ["" if c is None else str(c).strip() for c in row]
                if not any(raw):
                    continue
                # label with ALL value cells empty -> explicit blank marker
                if raw[0] and all(v == "" for v in raw[1:]):
                    lines.append(raw[0] + ":")
                    continue
                cells = [c for c in raw if c != ""]
                if not cells:
                    continue
                if len(cells) == 1:
                    lines.append(cells[0])
                else:
                    val = cells[1].split(" | ")[0]   # xlsx addresses follow " | "
                    lines.append(f"{cells[0]}: {val}")
                    for extra in cells[2:]:
                        lines.append(extra.split(" | ")[0])
        return (lines, "ok") if lines else ([], "empty")
    except Exception:
        return [], "corrupt"


def _docx_lines(data: bytes) -> tuple[list[str], str]:
    try:
        import docx  # python-docx
        document = docx.Document(io.BytesIO(data))
        lines: list[str] = []

        def table_rows(t):
            out = []
            for row in t.rows:
                raw = [c.text.strip() for c in row.cells]
                if not any(raw):
                    continue
                if raw[0] and all(v == "" for v in raw[1:]):
                    out.append(raw[0] + ":")
                    continue
                cells = [c for c in raw if c]
                if not cells:
                    continue
                if len(cells) == 1:
                    out.append(cells[0])
                else:
                    val_lines = [ln.strip() for ln in cells[1].splitlines() if ln.strip()]
                    if val_lines:
                        out.append(f"{cells[0]}: {val_lines[0]}")
                        out.extend(val_lines[1:])
                    for extra in cells[2:]:
                        out.append(extra)
            return out

        # interleave in document order where possible
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        body = document.element.body
        for child in body.iterchildren():
            if child.tag.endswith("}p"):
                p = Paragraph(child, document)
                if p.text.strip():
                    lines.append(p.text.strip())
            elif child.tag.endswith("}tbl"):
                t = Table(child, document)
                lines.extend(table_rows(t))
        return (lines, "ok") if lines else ([], "empty")
    except Exception:
        return [], "corrupt"


# ---------- document kind detection (wrong_doc_type guard) ----------
KIND_MARKERS = [
    ("commercial_invoice", [r"commercial\s+invoice"]),
    ("packing_list", [r"packing\s+list", r"packing\s+note"]),
    ("certificate_of_origin", [r"certificate\s+of\s+origin"]),
    ("invoice", [r"\binvoice\b"]),
    ("bill_of_lading", [r"bill\s+of\s+lading", r"\bb\s*/\s*l\b", r"\bbl\b"]),
    ("shipping_instruction", [r"shipping\s+instruction", r"\bsi\b"]),
]


def doc_kind(lines: list[str]) -> str:
    head = "\n".join(lines[:12]).lower()
    for kind, patterns in KIND_MARKERS:
        for p in patterns:
            if __import__("re").search(p, head):
                return kind
    return "unknown"


NOT_BL_KINDS = {"commercial_invoice", "packing_list", "certificate_of_origin"}
